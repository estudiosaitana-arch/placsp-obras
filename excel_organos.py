#!/usr/bin/env python3
"""Genera el listado Excel de órganos a partir de docs/data (ejecutar tras 'python placsp.py construir').
Uso: python excel_organos.py [salida.xlsx]"""
import json,glob,sys,os
os.chdir(os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
meta=json.load(open('docs/data/indice.json'))['meta']
import datetime as _dt
ANIOS=[str(a) for a in range(2023,_dt.date.today().year+1)]
rows=[]
for p in glob.glob('docs/data/o/*.json'):
    o=json.load(open(p,encoding='utf-8')); s=o['resumen']; pa=o['por_anio']; pr=o['procedimientos']
    n=s['expedientes'] or 0
    J=o['por_juicio_valor']['Con juicio de valor']; A=o['por_juicio_valor']['Solo criterios automáticos']
    ppr=lambda *ks: (sum(pr.get(k,0) for k in ks)/n) if n else None
    f=lambda v: None if v is None else v/100
    top='; '.join(f"{a['nombre']} ({a['n']})" for a in o['adjudicatarios'][:3])
    rows.append([o['nombre'],o['tipo'],o['provincia'],o['ciudad'],o['dependencia'],o['nif'],o['dir3'],n]
      +[(pa.get(a) or {}).get('n',0) for a in ANIOS]+[s['en_plazo'],
      s['pbl_total'],s['pbl_mediana'],s['adjudicado_total'],s['lotes_adjudicados'],
      f(s['baja_media']),f(s['baja_mediana']),f(s['baja_p75']),s['ofertas_media'],s['ofertas_mediana'],
      f(s['pct_desierto']),f(s['juicio_valor_medio']),
      f(s['pct_con_jv']),s['exp_con_jv'],s['exp_solo_auto'],s['exp_sin_criterios'],
      f(J['baja_media']),f(A['baja_media']),f(J['baja_mediana']),f(A['baja_mediana']),J['ofertas_media'],A['ofertas_media'],
      f(J['pct_desierto']),f(A['pct_desierto']),f(J['peso_precio_medio']),
      ppr('Abierto'),ppr('Abierto simplificado'),ppr('Negociado sin publicidad','Negociado con publicidad'),
      s['dias_fin_ofertas_a_adjudicacion'],f(s['pct_pyme_adjudicataria']),top,s['ultima_publicacion'] or None,
      o['email'],o['perfil'],f"https://estudiosaitana-arch.github.io/placsp-obras/#o={o['id']}",o['id']])
rows.sort(key=lambda r:(-r[7],-(r[13] or 0)))
H=[("Órgano de contratación",46,None),("Tipo de administración",24,None),("Provincia",16,None),("Ciudad",16,None),("Dependencia",50,None),("NIF",12,None),("DIR3",12,None),
("Expedientes obra (total)",12,'#,##0')]+[(f"Exp. {a}",9,'#,##0') for a in ANIOS]+[("En plazo ahora",9,'#,##0'),
("PBL total sin IVA (€)",16,'#,##0'),("PBL mediano (€)",14,'#,##0'),("Adjudicado total sin IVA (€)",16,'#,##0'),("Lotes adjudicados",10,'#,##0'),
("Baja media",9,'0.0%'),("Baja mediana",9,'0.0%'),("Baja P75",9,'0.0%'),("Ofertas por lote (media)",10,'0.0'),("Ofertas por lote (mediana)",10,'0.0'),
("% lotes desiertos",9,'0.0%'),("Peso medio juicio de valor",10,'0.0%'),
("% exp. con juicio de valor",10,'0.0%'),("Exp. con JV",9,'#,##0'),("Exp. solo automáticos",10,'#,##0'),("Exp. sin datos de criterios",10,'#,##0'),
("Baja media · con JV",10,'0.0%'),("Baja media · solo autom.",10,'0.0%'),("Baja mediana · con JV",10,'0.0%'),("Baja mediana · solo autom.",10,'0.0%'),
("Ofertas/lote · con JV",10,'0.0'),("Ofertas/lote · solo autom.",10,'0.0'),("% desiertos · con JV",10,'0.0%'),("% desiertos · solo autom.",10,'0.0%'),
("Peso precio medio · con JV",10,'0.0%'),
("% Abierto",9,'0%'),("% Abierto simplificado",10,'0%'),("% Negociado",9,'0%'),
("Días fin ofertas → adjudicación (mediana)",12,'#,##0'),("% adjudicado a pyme",9,'0.0%'),("Principales adjudicatarios (nº adj.)",60,None),("Último anuncio",11,None),
("Email contratación",28,None),("Perfil del contratante",14,None),("Ficha web",10,None),("ID órgano",16,None)]
wb=Workbook(); ws=wb.active; ws.title="Órganos"
F=Font(name="Arial",size=9); FB=Font(name="Arial",size=9,bold=True,color="FFFFFF"); FL=Font(name="Arial",size=9,color="0563C1",underline="single")
ws.append([h[0] for h in H])
for c,(h,w,fmt) in enumerate(H,1):
    cell=ws.cell(1,c); cell.font=FB; cell.fill=PatternFill("solid",fgColor="7A2E12"); cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
    ws.column_dimensions[get_column_letter(c)].width=w
ws.row_dimensions[1].height=48
for r in rows: ws.append(r)
iperf=[h[0] for h in H].index("Perfil del contratante")+1; ific=iperf+1
for row in ws.iter_rows(min_row=2):
    for c,cell in enumerate(row,1):
        cell.font=F
        if H[c-1][2]: cell.number_format=H[c-1][2]
    for col,txt in ((iperf,"Perfil"),(ific,"Ver ficha")):
        cell=row[col-1]
        if cell.value: cell.hyperlink=cell.value; cell.value=txt; cell.font=FL
ws.freeze_panes="B2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(H))}{ws.max_row}"
lg=wb.create_sheet("Leyenda")
L=[("Listado de órganos de contratación · licitaciones de obra",None),
("Fuente","Datos abiertos de la Plataforma de Contratación del Sector Público: sindicación 643 (PLACSP, sin menores) y 1044 (plataformas autonómicas agregadas, sin menores)."),
("Periodo","Expedientes de obra (tipo de contrato = Obras) publicados o actualizados en los meses cargados ("+min(p.split('_')[-1] for p in meta['periodos'])+" a "+meta['generado'][:10]+")."),
("Generado",meta['generado'].replace('T',' ')+" UTC · "+f"{meta['organos']:,} órganos · {meta['expedientes']:,} expedientes".replace(',', '.')),
("",None),
("Expedientes","Expedientes no anulados del órgano. El año se asigna por la fecha del anuncio de licitación (o, en su defecto, por la fecha fin de presentación de ofertas). Expedientes anteriores a 2023 que se actualizaron en el periodo cuentan en el total pero no en las columnas por año."),
("En plazo ahora","Expedientes en estado «En plazo» (PUB) en la última actualización."),
("PBL","Presupuesto base de licitación sin IVA del expediente."),
("Adjudicado total","Suma de importes de adjudicación sin IVA de los lotes adjudicados (excluye desiertos, desistimientos y renuncias)."),
("Baja","1 − importe adjudicado sin IVA / PBL sin IVA del lote. Se descartan valores fuera del rango 0–90 %. P75: percentil 75."),
("Ofertas por lote","Número de ofertas recibidas por lote adjudicado, según lo publicado por el órgano."),
("% lotes desiertos","Lotes con resultado desierto sobre el total de lotes con resultado publicado."),
("Peso juicio de valor","Peso medio de los criterios sujetos a juicio de valor (SUBJ) sobre la ponderación total, en los expedientes que publican ponderaciones."),
("Juicio de valor (JV)","Clasificación por el tipo de cada criterio de adjudicación publicado (cbc:AwardingCriteriaTypeCode): «con JV» si hay algún criterio SUBJ con peso > 0; «solo automáticos» si todos son OBJ; «sin datos» si el expediente no publica criterios (lo habitual en plataformas autonómicas agregadas). El % con JV se calcula sobre los expedientes que publican criterios."),
("Indicadores · con JV / solo autom.","Baja, ofertas por lote y % de desiertos calculados por separado para cada grupo. Peso precio medio · con JV: peso medio del criterio precio en los expedientes con juicio de valor."),
("% procedimiento","Proporción de expedientes por procedimiento sobre el total del órgano."),
("Días fin ofertas → adjudicación","Mediana de días entre la fecha límite de ofertas y la fecha de adjudicación (se descartan valores fuera de 0–730)."),
("Principales adjudicatarios","Tres empresas con más adjudicaciones de obra del órgano; entre paréntesis, número de lotes."),
("Tipo de administración","«Código N»: código de la plataforma agregada sin equivalencia confirmada. «Sin clasificar»: el órgano no informa el tipo."),
("Provincia","Por el código postal del órgano; si no lo publica, provincia donde se ubican la mayoría de sus obras."),
("",None),
("Limitaciones","Los indicadores dependen de lo que cada órgano publica: hay expedientes sin importe de adjudicación, sin número de ofertas o sin ponderación de criterios. Un mismo organismo puede figurar con varios perfiles (p. ej. Alcaldía y Junta de Gobierno). Las celdas vacías indican que no hay datos suficientes."),
]
for a,b in L: lg.append([a,b])
lg.column_dimensions['A'].width=30; lg.column_dimensions['B'].width=110
for row in lg.iter_rows():
    for cell in row: cell.font=Font(name="Arial",size=10,bold=(cell.column==1)); cell.alignment=Alignment(wrap_text=True,vertical="top")
lg['A1'].font=Font(name="Arial",size=13,bold=True)
salida=sys.argv[1] if len(sys.argv)>1 else 'data/organos_contratacion_obras_PLACSP.xlsx'
os.makedirs(os.path.dirname(salida) or '.',exist_ok=True)
wb.save(salida); print(f'{len(rows)} órganos -> {salida}')
